"""
================================================================================
Module:     helpers/biesse_converter.py
Project:    Biesse2Csv
Versie:     1.2.0
Datum:      2026-08-25
Auteur:     Bart Bossuyt

Doel:       Converteren van Biesse-machinebestanden naar CSV en XLSX.
            Leest machine-, operator- en EVENTOM-gegevens uit en schrijft
            deze zowel naar een Power BI-vriendelijk CSV-bestand als naar
            een direct in Excel bruikbaar XLSX-bestand.

Wijzigingen:
  v1.0.0 (2026-08-24)  Initiële versie — uitlezen van Biesse metadata,
                       operatorgegevens en EVENTOM-records naar CSV.
  v1.1.0 (2026-08-24)  XLSX-export toegevoegd naast CSV. Excel-bestand krijgt
                       tabelopmaak, filters, vastgezette header, automatische
                       kolombreedtes en correcte datum-/tijdnotatie.
  v1.1.1 (2026-08-24)  Foutieve openpyxl-import voor Table en TableStyleInfo
                       gecorrigeerd.
  v1.2.0 (2026-08-25)  XLSX uitgebreid met tabblad Samenvatting. Data wordt
                       gegroepeerd op Operator, Event_code,
                       Event_omschrijving en Event_type. Per groep worden
                       minimum- en maximumtijden, totale Duur_seconden en
                       Totale_duur opgenomen. Onderaan wordt een eindtotaal
                       toegevoegd.
================================================================================
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


@dataclass(frozen=True)
class ConversionResult:
    """Resultaat van één Biesse-conversie."""

    csv_file: Path
    xlsx_file: Path


class BiesseConverter:
    """Converteert Biesse-machinebestanden naar CSV en XLSX."""

    CSV_HEADERS = [
        "Machine",
        "Serienummer",
        "Frequentie",
        "Machine_start_datum",
        "Machine_start_tijd",
        "Machine_stop_datum",
        "Machine_stop_tijd",
        "Operator",
        "Operator_nummer",
        "Operator_shift",
        "Operator_start_datum",
        "Operator_start_tijd",
        "Operator_extra",
        "Event_code",
        "Event_omschrijving",
        "Duur_seconden",
        "Event_type",
        "Event_start_tijd",
        "Event_eind_tijd",
        "Operator_eind_datum",
        "Operator_eind_tijd",
        "Totale_duur",
    ]

    SUMMARY_HEADERS = [
        "Operator",
        "Event_code",
        "Event_omschrijving",
        "Event_type",
        "Operator_start_tijd",
        "Operator_stop_tijd",
        "Event_start_tijd",
        "Event_eind_tijd",
        "Totaal_Duur_seconden",
        "Totale_duur",
    ]

    DATE_COLUMNS = {
        "Machine_start_datum",
        "Machine_stop_datum",
        "Operator_start_datum",
        "Operator_eind_datum",
    }

    TIME_COLUMNS = {
        "Machine_start_tijd",
        "Machine_stop_tijd",
        "Operator_start_tijd",
        "Event_start_tijd",
        "Event_eind_tijd",
        "Operator_eind_tijd",
    }

    INTEGER_COLUMNS = {
        "Operator_nummer",
        "Operator_shift",
        "Duur_seconden",
        "Totale_duur",
    }

    SUMMARY_TIME_COLUMNS = {
        "Operator_start_tijd",
        "Operator_stop_tijd",
        "Event_start_tijd",
        "Event_eind_tijd",
    }

    SUMMARY_INTEGER_COLUMNS = {
        "Totaal_Duur_seconden",
        "Totale_duur",
    }

    def __init__(
        self,
        delimiter: str = ",",
        encoding: str = "utf-8-sig",
    ) -> None:
        """
        Initialiseer de converter.

        Args:
            delimiter:
                Scheidingsteken voor het CSV-bestand.

            encoding:
                Tekencodering voor het CSV-bestand.
        """
        self.delimiter = delimiter
        self.encoding = encoding

    # ==========================================================================
    # Publieke conversie
    # ==========================================================================

    def convert_file(
        self,
        source_file: Path | str,
        output_file: Path | str,
        *,
        overwrite: bool = False,
    ) -> ConversionResult:
        """
        Converteer één Biesse-bestand naar CSV én XLSX.

        Args:
            source_file:
                Pad naar het bronbestand.

            output_file:
                Gewenst pad van het CSV-bestand.
                Het XLSX-bestand krijgt automatisch dezelfde naam.

            overwrite:
                Indien False worden bestaande outputbestanden niet
                overschreven.

        Returns:
            ConversionResult met het CSV- en XLSX-pad.

        Raises:
            FileNotFoundError:
                Wanneer het bronbestand niet bestaat.

            FileExistsError:
                Wanneer CSV of XLSX al bestaat en overwrite=False.

            ValueError:
                Wanneer geen EVENTOM-records gevonden worden.
        """
        source_path = Path(source_file)
        csv_path = Path(output_file)

        if csv_path.suffix.lower() != ".csv":
            csv_path = csv_path.with_suffix(".csv")

        xlsx_path = csv_path.with_suffix(".xlsx")

        if not source_path.is_file():
            raise FileNotFoundError(
                f"Bronbestand bestaat niet: {source_path}"
            )

        if not overwrite:
            existing_files = []

            if csv_path.exists():
                existing_files.append(csv_path)

            if xlsx_path.exists():
                existing_files.append(xlsx_path)

            if existing_files:
                filenames = ", ".join(
                    path.name
                    for path in existing_files
                )

                raise FileExistsError(
                    f"Outputbestand bestaat reeds: {filenames}"
                )

        text = self._read_source(
            source_path
        )

        metadata = self._extract_metadata(
            text
        )

        operator_start = self._extract_operator_start(
            text
        )

        operator_end = self._extract_operator_end(
            text
        )

        events = self._extract_events(
            text
        )

        if not events:
            raise ValueError(
                f"Geen EVENTOM-records gevonden in: {source_path.name}"
            )

        rows = self._build_rows(
            metadata=metadata,
            operator_start=operator_start,
            operator_end=operator_end,
            events=events,
        )

        csv_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        self._write_csv(
            output_path=csv_path,
            rows=rows,
        )

        self._write_xlsx(
            output_path=xlsx_path,
            rows=rows,
        )

        return ConversionResult(
            csv_file=csv_path,
            xlsx_file=xlsx_path,
        )

    def convert_folder(
        self,
        input_folder: Path | str,
        output_folder: Path | str,
        *,
        overwrite: bool = False,
    ) -> list[ConversionResult]:
        """
        Converteer alle geschikte Biesse-bestanden uit een map.

        Returns:
            Lijst met ConversionResult-objecten.
        """
        input_path = Path(
            input_folder
        )

        output_path = Path(
            output_folder
        )

        if not input_path.is_dir():
            raise NotADirectoryError(
                f"Ongeldige bronmap: {input_path}"
            )

        output_path.mkdir(
            parents=True,
            exist_ok=True,
        )

        converted_files: list[ConversionResult] = []

        for source_file in sorted(
            input_path.iterdir()
        ):
            if not source_file.is_file():
                continue

            if not self.is_biesse_file(
                source_file
            ):
                continue

            csv_file = (
                output_path
                / self.build_output_name(
                    source_file
                )
            )

            result = self.convert_file(
                source_file=source_file,
                output_file=csv_file,
                overwrite=overwrite,
            )

            converted_files.append(
                result
            )

        return converted_files

    # ==========================================================================
    # Bestandsnamen
    # ==========================================================================

    @staticmethod
    def is_biesse_file(
        file_path: Path | str,
    ) -> bool:
        """
        Controleer of een bestand als Biesse-bronbestand wordt beschouwd.

        Voorbeelden:
            M2608181.520 -> True
            machine.001  -> True
            test.csv     -> False
        """
        path = Path(
            file_path
        )

        suffix = path.suffix.lstrip(".")

        return bool(
            suffix
            and suffix.isdigit()
        )

    @staticmethod
    def build_output_name(
        source_file: Path | str,
    ) -> str:
        """
        Bouw de CSV-bestandsnaam.

        Voorbeeld:
            M2608181.520
            ->
            M2608181_520.csv
        """
        path = Path(
            source_file
        )

        extension = path.suffix.lstrip(".")

        return (
            f"{path.stem}_{extension}.csv"
        )

    @staticmethod
    def build_xlsx_output_name(
        source_file: Path | str,
    ) -> str:
        """
        Bouw de XLSX-bestandsnaam.

        Voorbeeld:
            M2608181.520
            ->
            M2608181_520.xlsx
        """
        path = Path(
            source_file
        )

        extension = path.suffix.lstrip(".")

        return (
            f"{path.stem}_{extension}.xlsx"
        )

    # ==========================================================================
    # Bronbestand lezen
    # ==========================================================================

    @staticmethod
    def _read_source(
        source_path: Path,
    ) -> str:
        """Lees het Biesse-bestand als Windows-1252."""

        return source_path.read_bytes().decode(
            "cp1252",
            errors="replace",
        )

    # ==========================================================================
    # Metadata
    # ==========================================================================

    @staticmethod
    def _extract_field(
        text: str,
        field_name: str,
    ) -> str:
        """Lees een ^VELD=waarde uit het bronbestand."""

        match = re.search(
            rf"\^{re.escape(field_name)}=([^\^\r\n\t]*)",
            text,
        )

        if not match:
            return ""

        return match.group(1).strip()

    def _extract_metadata(
        self,
        text: str,
    ) -> dict[str, str]:
        """Lees algemene machinegegevens."""

        return {
            "machine": self._extract_field(
                text,
                "NOME",
            ),
            "serial_number": self._extract_field(
                text,
                "MATRICOLA",
            ),
            "frequency": self._extract_field(
                text,
                "SCADMAC",
            ),
            "machine_start_time": self._extract_field(
                text,
                "OSTARTMAC",
            ),
            "machine_start_date": self._extract_field(
                text,
                "DSTARTMAC",
            ),
            "machine_stop_time": self._extract_field(
                text,
                "OSTOPMAC",
            ),
            "machine_stop_date": self._extract_field(
                text,
                "DSTOPMAC",
            ),
        }

    # ==========================================================================
    # Record parser
    # ==========================================================================

    @staticmethod
    def _parse_tab_row(
        row_text: str,
    ) -> list[str]:
        """Parse een tab-gescheiden Biesse-regel."""

        try:
            return next(
                csv.reader(
                    [row_text],
                    delimiter="\t",
                    quotechar='"',
                )
            )

        except (
            csv.Error,
            StopIteration,
        ):
            return []

    def _extract_operator_start(
        self,
        text: str,
    ) -> list[str]:
        """Lees de operator-startregel."""

        match = re.search(
            r"\^OPERATOREINI=\s*\r?\n\s*(\".*\")",
            text,
        )

        if not match:
            return []

        return self._parse_tab_row(
            match.group(1)
        )

    def _extract_operator_end(
        self,
        text: str,
    ) -> list[str]:
        """Lees de operator-eindregel."""

        match = re.search(
            r"\^OPERATOREFIN=\s*\r?\n\s*(\".*\")",
            text,
        )

        if not match:
            return []

        return self._parse_tab_row(
            match.group(1)
        )

    def _extract_events(
        self,
        text: str,
    ) -> list[list[str]]:
        """Lees alle EVENTOM-records."""

        pattern = re.compile(
            r"\^EVENTOM=\s*\r?\n\s*(\".*\")"
        )

        events: list[list[str]] = []

        for match in pattern.finditer(
            text
        ):
            values = self._parse_tab_row(
                match.group(1)
            )

            if len(values) < 6:
                continue

            events.append(
                values[:6]
            )

        return events

    # ==========================================================================
    # CSV-rijen
    # ==========================================================================

    @staticmethod
    def _safe_get(
        values: list[str],
        index: int,
    ) -> str:
        """Lees veilig een lijstwaarde."""

        if index < len(values):
            return values[index]

        return ""

    def _build_rows(
        self,
        metadata: dict[str, str],
        operator_start: list[str],
        operator_end: list[str],
        events: list[list[str]],
    ) -> list[list[str]]:
        """Bouw de uiteindelijke data-rijen."""

        rows: list[list[str]] = []

        for event in events:
            row = [
                metadata["machine"],
                metadata["serial_number"],
                metadata["frequency"],
                metadata["machine_start_date"],
                metadata["machine_start_time"],
                metadata["machine_stop_date"],
                metadata["machine_stop_time"],

                self._safe_get(
                    operator_start,
                    0,
                ),
                self._safe_get(
                    operator_start,
                    1,
                ),
                self._safe_get(
                    operator_start,
                    2,
                ),
                self._safe_get(
                    operator_start,
                    4,
                ),
                self._safe_get(
                    operator_start,
                    3,
                ),
                self._safe_get(
                    operator_start,
                    5,
                ),

                self._safe_get(
                    event,
                    0,
                ),
                self._safe_get(
                    event,
                    1,
                ),
                self._safe_get(
                    event,
                    2,
                ),
                self._safe_get(
                    event,
                    3,
                ),
                self._safe_get(
                    event,
                    4,
                ),
                self._safe_get(
                    event,
                    5,
                ),

                self._safe_get(
                    operator_end,
                    1,
                ),
                self._safe_get(
                    operator_end,
                    0,
                ),
                self._safe_get(
                    operator_end,
                    2,
                ),
            ]

            rows.append(
                row
            )

        return rows

    # ==========================================================================
    # Samenvatting
    # ==========================================================================

    def _build_summary_rows(
        self,
        rows: list[list[str]],
    ) -> list[list]:
        """
        Groepeer detailrijen voor het Excel-tabblad Samenvatting.

        Groepering:
            - Operator
            - Event_code
            - Event_omschrijving
            - Event_type

        Berekeningen per groep:
            - Operator_start_tijd = minimum
            - Operator_stop_tijd = maximum Operator_eind_tijd
            - Event_start_tijd = minimum
            - Event_eind_tijd = maximum
            - Totaal_Duur_seconden = som Duur_seconden
            - Totale_duur = maximum Totale_duur binnen de groep

        Totale_duur wordt bewust NIET opgeteld, omdat deze waarde in de
        detaildata op iedere EVENTOM-regel terugkomt.
        """

        header_index = {
            header: index
            for index, header in enumerate(
                self.CSV_HEADERS
            )
        }

        grouped: dict[
            tuple[str, str, str, str],
            dict,
        ] = {}

        for row in rows:
            operator = str(
                row[
                    header_index["Operator"]
                ]
            ).strip()

            event_code = str(
                row[
                    header_index["Event_code"]
                ]
            ).strip()

            event_description = str(
                row[
                    header_index["Event_omschrijving"]
                ]
            ).strip()

            event_type = str(
                row[
                    header_index["Event_type"]
                ]
            ).strip()

            key = (
                operator,
                event_code,
                event_description,
                event_type,
            )

            operator_start = self._parse_time_value(
                row[
                    header_index["Operator_start_tijd"]
                ]
            )

            operator_stop = self._parse_time_value(
                row[
                    header_index["Operator_eind_tijd"]
                ]
            )

            event_start = self._parse_time_value(
                row[
                    header_index["Event_start_tijd"]
                ]
            )

            event_end = self._parse_time_value(
                row[
                    header_index["Event_eind_tijd"]
                ]
            )

            duration = self._parse_int_value(
                row[
                    header_index["Duur_seconden"]
                ]
            )

            total_duration = self._parse_int_value(
                row[
                    header_index["Totale_duur"]
                ]
            )

            if key not in grouped:
                grouped[key] = {
                    "operator_start": operator_start,
                    "operator_stop": operator_stop,
                    "event_start": event_start,
                    "event_end": event_end,
                    "duration": duration,
                    "total_duration": total_duration,
                }

                continue

            group = grouped[key]

            group["operator_start"] = self._min_time(
                group["operator_start"],
                operator_start,
            )

            group["operator_stop"] = self._max_time(
                group["operator_stop"],
                operator_stop,
            )

            group["event_start"] = self._min_time(
                group["event_start"],
                event_start,
            )

            group["event_end"] = self._max_time(
                group["event_end"],
                event_end,
            )

            group["duration"] += duration

            group["total_duration"] = max(
                group["total_duration"],
                total_duration,
            )

        summary_rows: list[list] = []

        for key, values in sorted(
            grouped.items(),
            key=lambda item: (
                item[0][0],
                item[0][1],
                item[0][2],
                item[0][3],
            ),
        ):
            operator, event_code, event_description, event_type = key

            summary_rows.append(
                [
                    operator,
                    event_code,
                    event_description,
                    event_type,
                    values["operator_start"],
                    values["operator_stop"],
                    values["event_start"],
                    values["event_end"],
                    values["duration"],
                    values["total_duration"],
                ]
            )

        return summary_rows

    # ==========================================================================
    # CSV schrijven
    # ==========================================================================

    def _write_csv(
        self,
        output_path: Path,
        rows: list[list[str]],
    ) -> None:
        """Schrijf de data naar CSV."""

        with output_path.open(
            "w",
            newline="",
            encoding=self.encoding,
        ) as file:
            writer = csv.writer(
                file,
                delimiter=self.delimiter,
                quotechar='"',
                quoting=csv.QUOTE_MINIMAL,
            )

            writer.writerow(
                self.CSV_HEADERS
            )

            writer.writerows(
                rows
            )

    # ==========================================================================
    # XLSX schrijven
    # ==========================================================================

    def _write_xlsx(
        self,
        output_path: Path,
        rows: list[list[str]],
    ) -> None:
        """
        Schrijf gegevens naar een echt Excel XLSX-bestand.

        Het bestand bevat:
            - tabblad Biesse data met alle detailregels;
            - tabblad Samenvatting met gegroepeerde gegevens;
            - filterknoppen;
            - tabelopmaak;
            - vastgezette headers;
            - correcte datum- en tijdwaarden;
            - aangepaste kolombreedtes;
            - eindtotaal op het tabblad Samenvatting.
        """

        workbook = Workbook()

        # ======================================================================
        # Tabblad 1: Biesse data
        # ======================================================================

        worksheet = workbook.active
        worksheet.title = "Biesse data"

        worksheet.append(
            self.CSV_HEADERS
        )

        for source_row in rows:
            excel_row = []

            for header, value in zip(
                self.CSV_HEADERS,
                source_row,
            ):
                excel_row.append(
                    self._convert_excel_value(
                        header,
                        value,
                    )
                )

            worksheet.append(
                excel_row
            )

        worksheet.freeze_panes = "A2"

        if rows:
            last_row = worksheet.max_row
            last_column = worksheet.max_column

            table_range = (
                f"A1:"
                f"{get_column_letter(last_column)}"
                f"{last_row}"
            )

            table = Table(
                displayName="BiesseData",
                ref=table_range,
            )

            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = style

            worksheet.add_table(
                table
            )

        self._format_detail_worksheet(
            worksheet
        )

        # ======================================================================
        # Tabblad 2: Samenvatting
        # ======================================================================

        summary_worksheet = workbook.create_sheet(
            title="Samenvatting"
        )

        summary_rows = self._build_summary_rows(
            rows
        )

        summary_worksheet.append(
            self.SUMMARY_HEADERS
        )

        for summary_row in summary_rows:
            summary_worksheet.append(
                summary_row
            )

        summary_worksheet.freeze_panes = "A2"

        # ----------------------------------------------------------------------
        # Excel Table
        # ----------------------------------------------------------------------

        if summary_rows:
            last_data_row = (
                len(summary_rows) + 1
            )

            last_column = len(
                self.SUMMARY_HEADERS
            )

            table_range = (
                f"A1:"
                f"{get_column_letter(last_column)}"
                f"{last_data_row}"
            )

            summary_table = Table(
                displayName="BiesseSamenvatting",
                ref=table_range,
            )

            summary_style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            summary_table.tableStyleInfo = summary_style

            summary_worksheet.add_table(
                summary_table
            )

        # ----------------------------------------------------------------------
        # Eindtotaal
        # ----------------------------------------------------------------------

        if summary_rows:
            total_row_number = (
                len(summary_rows) + 3
            )

            summary_worksheet.cell(
                row=total_row_number,
                column=1,
                value="Eindtotaal",
            )

            total_duration_seconds = sum(
                row[8]
                for row in summary_rows
                if isinstance(
                    row[8],
                    int,
                )
            )

            summary_worksheet.cell(
                row=total_row_number,
                column=9,
                value=total_duration_seconds,
            )

            # Totale_duur is operator-/sessieduur en komt in elke
            # detailregel terug. Daarom niet optellen per eventgroep.
            #
            # Voor het eindtotaal wordt per operator slechts de hoogste
            # Totale_duur meegenomen.
            operator_totals: dict[str, int] = {}

            for summary_row in summary_rows:
                operator = str(
                    summary_row[0]
                )

                total_duration = (
                    summary_row[9]
                    if isinstance(
                        summary_row[9],
                        int,
                    )
                    else 0
                )

                operator_totals[operator] = max(
                    operator_totals.get(
                        operator,
                        0,
                    ),
                    total_duration,
                )

            summary_worksheet.cell(
                row=total_row_number,
                column=10,
                value=sum(
                    operator_totals.values()
                ),
            )

            for cell in summary_worksheet[
                total_row_number
            ]:
                cell.font = Font(
                    bold=True
                )

        self._format_summary_worksheet(
            summary_worksheet
        )

        # ======================================================================
        # Opslaan
        # ======================================================================

        workbook.save(
            output_path
        )

    # ==========================================================================
    # Excel formattering
    # ==========================================================================

    def _format_detail_worksheet(
        self,
        worksheet,
    ) -> None:
        """Formatteer het detailtabblad."""

        header_map = {
            cell.value: cell.column
            for cell in worksheet[1]
        }

        for header in self.DATE_COLUMNS:
            column = header_map.get(
                header
            )

            if column is None:
                continue

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column,
                ).number_format = "dd/mm/yyyy"

        for header in self.TIME_COLUMNS:
            column = header_map.get(
                header
            )

            if column is None:
                continue

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column,
                ).number_format = "hh:mm:ss"

        for header in self.INTEGER_COLUMNS:
            column = header_map.get(
                header
            )

            if column is None:
                continue

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column,
                ).number_format = "0"

        self._style_header(
            worksheet
        )

        self._set_excel_column_widths(
            worksheet
        )

    def _format_summary_worksheet(
        self,
        worksheet,
    ) -> None:
        """Formatteer het tabblad Samenvatting."""

        header_map = {
            cell.value: cell.column
            for cell in worksheet[1]
        }

        for header in self.SUMMARY_TIME_COLUMNS:
            column = header_map.get(
                header
            )

            if column is None:
                continue

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column,
                ).number_format = "hh:mm:ss"

        for header in self.SUMMARY_INTEGER_COLUMNS:
            column = header_map.get(
                header
            )

            if column is None:
                continue

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column,
                ).number_format = "0"

        self._style_header(
            worksheet
        )

        self._set_excel_column_widths(
            worksheet
        )

    @staticmethod
    def _style_header(
        worksheet,
    ) -> None:
        """Maak de header vet en verticaal uitgelijnd."""

        for cell in worksheet[1]:
            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                vertical="center",
            )

    # ==========================================================================
    # Excel helpers
    # ==========================================================================

    def _convert_excel_value(
        self,
        header: str,
        value: str,
    ):
        """
        Converteer tekstwaarden naar geschikte Excel-datatypes.
        """

        if value is None:
            return ""

        value = str(
            value
        ).strip()

        if not value:
            return ""

        if header in self.DATE_COLUMNS:
            try:
                return datetime.strptime(
                    value,
                    "%d/%m/%Y",
                ).date()

            except ValueError:
                return value

        if header in self.TIME_COLUMNS:
            try:
                return datetime.strptime(
                    value,
                    "%H:%M:%S",
                ).time()

            except ValueError:
                return value

        if header in self.INTEGER_COLUMNS:
            try:
                return int(
                    value
                )

            except ValueError:
                return value

        return value

    @staticmethod
    def _parse_time_value(
        value,
    ) -> time | None:
        """
        Converteer een tijdwaarde naar datetime.time.

        Ongeldige of lege waarden geven None terug.
        """

        if value is None:
            return None

        if isinstance(
            value,
            time,
        ):
            return value

        text = str(
            value
        ).strip()

        if not text:
            return None

        try:
            return datetime.strptime(
                text,
                "%H:%M:%S",
            ).time()

        except ValueError:
            return None

    @staticmethod
    def _parse_int_value(
        value,
    ) -> int:
        """
        Converteer een waarde veilig naar integer.

        Ongeldige of lege waarden worden 0.
        """

        if value is None:
            return 0

        try:
            return int(
                str(
                    value
                ).strip()
            )

        except (
            ValueError,
            TypeError,
        ):
            return 0

    @staticmethod
    def _min_time(
        current: time | None,
        new_value: time | None,
    ) -> time | None:
        """Geef de vroegste niet-lege tijd terug."""

        if current is None:
            return new_value

        if new_value is None:
            return current

        return min(
            current,
            new_value,
        )

    @staticmethod
    def _max_time(
        current: time | None,
        new_value: time | None,
    ) -> time | None:
        """Geef de laatste niet-lege tijd terug."""

        if current is None:
            return new_value

        if new_value is None:
            return current

        return max(
            current,
            new_value,
        )

    @staticmethod
    def _set_excel_column_widths(
        worksheet,
    ) -> None:
        """Bepaal leesbare kolombreedtes."""

        for column_cells in worksheet.columns:
            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                if cell.value is None:
                    continue

                text = str(
                    cell.value
                )

                max_length = max(
                    max_length,
                    len(text),
                )

            # Niet onbeperkt breed laten worden.
            width = min(
                max(
                    max_length + 2,
                    10,
                ),
                32,
            )

            worksheet.column_dimensions[
                column_letter
            ].width = width